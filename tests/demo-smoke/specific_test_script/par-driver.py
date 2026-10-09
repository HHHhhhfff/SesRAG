"""PAR_001/PAR_003/PAR_004：生成本次docx并调用真实解析函数。"""

import base64
import json
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

from backend.infrastructure.rag.parsers import parse_document

result = {}
with tempfile.TemporaryDirectory(prefix="par-fixtures-") as directory:
    root = Path(directory)
    sources = {
        "txt": (".txt", "中英UTF8保真\nSecond line\n末尾🙂"),
        "md": (".md", "# 标题\n**Markdown标记** 与 `代码`\n末尾"),
    }
    # PAR_001：Markdown保留标记；offset使用解析文本字符而非UTF8字节。
    for name, (extension, text) in sources.items():
        path = root / (name + extension)
        path.write_text(text, encoding="utf-8")
        parsed = parse_document(path, extension)
        result[name] = {
            "bytes": base64.b64encode(path.read_bytes()).decode(),
            "text": parsed.text,
            "expected": text,
            "source_name": parsed.source_name,
            "position_fields": [
                field
                for field in ("positions", "locations", "start_offset", "end_offset")
                if hasattr(parsed, field)
            ],
        }
    for name in ("docx", "empty_docx", "table_only", "mixed_regions"):
        document = Document()
        expected = ""
        if name == "docx":
            document.add_paragraph("第一段中文🙂")
            document.add_paragraph("Second paragraph")
            expected = "第一段中文🙂\nSecond paragraph"
        if name == "table_only":
            document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_ONLY_SENTINEL"
        if name == "mixed_regions":
            document.add_paragraph("BODY_SENTINEL")
            document.sections[0].header.paragraphs[0].text = "HEADER_SENTINEL"
            document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_SENTINEL"
            document.paragraphs[0]._p.append(
                parse_xml(
                    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    "<w:pict><w:txbxContent><w:p><w:r><w:t>TEXTBOX_SENTINEL</w:t></w:r>"
                    "</w:p></w:txbxContent></w:pict></w:r>"
                )
            )
            expected = "BODY_SENTINEL"
        path = root / (name + ".docx")
        document.save(path)
        parsed = parse_document(path, ".docx")
        result[name] = {
            "bytes": base64.b64encode(path.read_bytes()).decode(),
            "text": parsed.text,
            "expected": expected,
            "source_name": parsed.source_name,
            "position_fields": [
                field
                for field in ("positions", "locations", "start_offset", "end_offset")
                if hasattr(parsed, field)
            ],
            "regions": {
                region: sentinel in parsed.text
                for region, sentinel in {
                    "body": "BODY_SENTINEL",
                    "header": "HEADER_SENTINEL",
                    "table": "TABLE_SENTINEL",
                    "textbox": "TEXTBOX_SENTINEL",
                }.items()
            },
        }
print("SMOKE_JSON:" + json.dumps(result, ensure_ascii=False))
