"""IDX_001/IDX_002：Mock向量实际调用及本次三格式输入。"""

import base64
import io
import json
import math

from docx import Document

from backend.infrastructure.providers.mock import MockEmbeddingProvider

provider = MockEmbeddingProvider(384)
same = "IDX_001相同中英文 stable vector"
vectors = [provider.embed(same), provider.embed(same), provider.embed("IDX_001不同安全文本")]
document = Document()
document.add_paragraph("IDX_002第一段索引资料")
document.add_paragraph("Second paragraph for indexed evidence")
stream = io.BytesIO()
document.save(stream)
result = {
    "vectors": {
        "equal_repeated": vectors[0] == vectors[1],
        "dimensions": [len(vector) for vector in vectors],
        "all_finite": all(math.isfinite(value) for vector in vectors for value in vector),
        "norms": [math.sqrt(sum(value * value for value in vector)) for vector in vectors],
    },
    "fixtures": {
        "txt": {
            "bytes": base64.b64encode("IDX_002中文🙂\nUTF8 safe text".encode()).decode(),
            "text": "IDX_002中文🙂\nUTF8 safe text",
        },
        "md": {
            "bytes": base64.b64encode("# IDX_002标题\n**安全资料**".encode()).decode(),
            "text": "# IDX_002标题\n**安全资料**",
        },
        "docx": {
            "bytes": base64.b64encode(stream.getvalue()).decode(),
            "text": "IDX_002第一段索引资料\nSecond paragraph for indexed evidence",
        },
    },
}
print("SMOKE_JSON:" + json.dumps(result, ensure_ascii=False))
