"""PAR fixture兼容回归：保留真docx字节，限定空错误，其他异常原样传播。"""

import base64
import io
import json
import runpy
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from docx import Document

from backend.application.errors import PermanentTaskError, TransientTaskError
from backend.infrastructure.rag import parsers

DRIVER = Path(__file__).parents[1] / "specific_test_script" / "par-driver.py"


class ParFixtureDriver(unittest.TestCase):
    def run_driver(self, *, target=None, failure=None):
        original_parse = parsers.parse_document
        original_temp = tempfile.TemporaryDirectory
        saved_bytes, directories = {}, []

        def parse(path, extension):
            saved_bytes[path.stem] = path.read_bytes()
            if path.stem == target:
                raise failure
            return original_parse(path, extension)

        with original_temp(prefix="par-driver-test-") as owned:

            def temporary(**kwargs):
                result = original_temp(dir=owned, **kwargs)
                directories.append(Path(result.name))
                return result

            output = io.StringIO()
            try:
                with (
                    patch.object(parsers, "parse_document", parse),
                    patch.object(tempfile, "TemporaryDirectory", temporary),
                    redirect_stdout(output),
                ):
                    runpy.run_path(str(DRIVER))
            finally:
                # 成功和异常都必须让driver自己的TemporaryDirectory退出。
                self.assertTrue(directories)
                self.assertTrue(all(not directory.exists() for directory in directories))
            rows = [
                line for line in output.getvalue().splitlines() if line.startswith("SMOKE_JSON:")
            ]
            self.assertEqual(len(rows), 1)
            return json.loads(rows[0].removeprefix("SMOKE_JSON:")), saved_bytes

    def test_real_empty_fixtures_remain_uploadable_and_normals_keep_parser_results(self):
        try:
            result, saved = self.run_driver()
        except PermanentTaskError as error:
            self.fail(f"预期空fixture不应中止整份输出：{error.code}")
        for name in ("txt", "md", "docx", "mixed_regions"):
            with self.subTest(name=name):
                self.assertEqual(result[name]["text"], result[name]["expected"])
                self.assertTrue(result[name]["text"])
                self.assertIn("locations", result[name]["position_fields"])
                self.assertEqual(base64.b64decode(result[name]["bytes"]), saved[name])
        for name, tables in (("empty_docx", 0), ("table_only", 1)):
            with self.subTest(name=name):
                fixture = result[name]
                binary = base64.b64decode(fixture["bytes"])
                self.assertEqual(binary, saved[name])
                loaded = Document(io.BytesIO(binary))
                self.assertEqual(loaded.paragraphs, [])
                self.assertEqual(len(loaded.tables), tables)
                if tables:
                    self.assertEqual(loaded.tables[0].cell(0, 0).text, "TABLE_ONLY_SENTINEL")
                self.assertEqual(fixture["text"], "")
                self.assertEqual(fixture["parse_error_code"], "EMPTY_DOCUMENT")
                self.assertEqual(fixture["parse_status"], "expected_empty_error")
                self.assertEqual(
                    fixture["parse_result_source"],
                    "backend.infrastructure.rag.parsers.parse_document",
                )
                self.assertEqual(fixture["source_name"], name + ".docx")
                self.assertEqual(fixture["position_fields"], [])

    def test_empty_document_error_on_normal_fixtures_is_not_swallowed(self):
        for name in ("txt", "md", "docx", "mixed_regions"):
            with self.subTest(name=name):
                failure = PermanentTaskError("EMPTY_DOCUMENT", "正常fixture不应为空")
                with self.assertRaises(PermanentTaskError) as caught:
                    self.run_driver(target=name, failure=failure)
                self.assertIs(caught.exception, failure)

    def test_other_errors_on_expected_empty_fixtures_are_not_swallowed(self):
        for name in ("empty_docx", "table_only"):
            for failure in (
                PermanentTaskError("DOCUMENT_PARSE_FAILED", "结构错误"),
                TransientTaskError("DOCUMENT_STORAGE_FAILED", "存储错误"),
                RuntimeError("未知解析故障"),
            ):
                with self.subTest(name=name, failure=type(failure).__name__):
                    with self.assertRaises(type(failure)) as caught:
                        self.run_driver(target=name, failure=failure)
                    self.assertIs(caught.exception, failure)


if __name__ == "__main__":
    unittest.main()
