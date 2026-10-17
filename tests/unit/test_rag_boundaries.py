import asyncio
import json

import httpx
import pytest

from backend.application.errors import ApplicationError, PermanentTaskError
from backend.infrastructure.providers.openai_compatible import OpenAICompatibleProvider
from backend.infrastructure.rag.parsers import parse_document
from backend.infrastructure.rag.pgvector_store import search_chunks


@pytest.mark.parametrize(
    "extension,content,code",
    [
        (".txt", b"\xff", "DOCUMENT_ENCODING_INVALID"),
        (".docx", b"invalid", "DOCUMENT_PARSE_FAILED"),
        (".txt", b"", "EMPTY_DOCUMENT"),
    ],
)
def test_parser_rejects_permanent_bad_inputs(tmp_path, extension, content, code):
    path = tmp_path / ("secret-source" + extension)
    path.write_bytes(content)
    with pytest.raises(PermanentTaskError) as error:
        parse_document(path, extension)
    assert error.value.code == code
    assert "secret-source" not in str(error.value)


@pytest.mark.asyncio
async def test_parser_timeout_kills_and_reaps_actual_child(tmp_path, monkeypatch):
    import sys

    from backend.infrastructure.rag import parsers

    original = asyncio.create_subprocess_exec
    children = []

    async def spawn(*args, **kwargs):
        child = await original(sys.executable, "-c", "import time; time.sleep(60)", **kwargs)
        children.append(child)
        return child

    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    monkeypatch.setattr(parsers, "PARSE_TIMEOUT", 0.05)
    with pytest.raises(PermanentTaskError, match="解析超时"):
        await parsers.SubprocessParser(tmp_path).parse("source.txt", ".txt")
    assert len(children) == 1 and children[0].returncode is not None


@pytest.mark.parametrize(
    "values",
    [
        {"mock_provider": False},
        {
            "mock_provider": False,
            "model_base_url": "http://local/v1",
            "model_name": "model",
            "embedding_model_name": "embed",
        },
    ],
)
def test_real_provider_configuration_requires_all_values(values):
    from pydantic import ValidationError

    from backend.config import Settings

    with pytest.raises(ValidationError):
        Settings(_env_file=None, **values)


def test_parser_returns_line_and_paragraph_positions(tmp_path):
    from docx import Document

    path = tmp_path / "note.txt"
    path.write_text("first\nsecond", encoding="utf-8")
    parsed = parse_document(path, ".txt")
    assert [(p.kind, p.number, p.start_offset, p.end_offset) for p in parsed.locations] == [
        ("line", 1, 0, 5),
        ("line", 2, 6, 12),
    ]
    document = Document()
    document.add_paragraph("alpha")
    document.add_paragraph("beta")
    path = tmp_path / "note.docx"
    document.save(path)
    parsed = parse_document(path, ".docx")
    assert parsed.text == "alpha\nbeta"
    assert [(p.kind, p.number) for p in parsed.locations] == [("paragraph", 1), ("paragraph", 2)]


def test_docx_soft_break_is_one_paragraph_and_empty_paragraph_keeps_number(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("a\nb")
    document.add_paragraph("")
    document.add_paragraph("c")
    path = tmp_path / "soft.docx"
    document.save(path)
    parsed = parse_document(path, ".docx")
    assert parsed.text == "a\nb\n\nc"
    assert [(p.number, p.start_offset, p.end_offset) for p in parsed.locations] == [
        (1, 0, 3),
        (2, 4, 4),
        (3, 5, 6),
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "vector,k,code",
    [
        ([1.0] * 384, 0, "INVALID_TOP_K"),
        ([1.0], 1, "EMBEDDING_DIMENSION_MISMATCH"),
        ([float("nan")] * 384, 1, "EMBEDDING_INVALID"),
        ([1e308] * 384, 1, "EMBEDDING_INVALID"),
    ],
)
async def test_search_rejects_invalid_input_before_query(vector, k, code):
    with pytest.raises(ApplicationError) as error:
        await search_chunks(None, "owner", vector, k)
    assert error.value.code == code


def provider_with_response(monkeypatch, response):
    original = httpx.AsyncClient
    transport = httpx.MockTransport(lambda request: response(request))
    monkeypatch.setattr(
        httpx, "AsyncClient", lambda **kwargs: original(transport=transport, **kwargs)
    )
    return OpenAICompatibleProvider("http://local/v1", "secret", "model", "embed")


@pytest.mark.asyncio
async def test_provider_consumes_three_sse_deltas(monkeypatch):
    def response(request):
        assert json.loads(request.content)["stream"] is True
        payload = (
            "".join(
                "data: " + json.dumps({"choices": [{"delta": {"content": part}}]}) + "\n\n"
                for part in ["one", "two", "three"]
            )
            + "data: [DONE]\n\n"
        )
        return httpx.Response(200, text=payload, headers={"content-type": "text/event-stream"})

    provider = provider_with_response(monkeypatch, response)
    assert [part async for part in provider.stream_answer("private prompt", [])] == [
        "one",
        "two",
        "three",
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "status,retryable", [(400, False), (401, False), (403, False), (429, True), (500, True)]
)
async def test_provider_classifies_http_without_leaking_body(monkeypatch, status, retryable):
    provider = provider_with_response(
        monkeypatch, lambda _: httpx.Response(status, text="secret /private/path")
    )
    with pytest.raises(ApplicationError) as error:
        await provider.embed("private prompt")
    assert error.value.retryable is retryable
    assert error.value.code == f"PROVIDER_HTTP_{status}"
    assert "secret" not in str(error.value)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "body,code",
    [
        ("not json", "PROVIDER_INVALID_RESPONSE"),
        ('{"data":[{"embedding":[1]}]}', "EMBEDDING_DIMENSION_MISMATCH"),
        (json.dumps({"data": [{"embedding": [float("inf")] * 384}]}), "EMBEDDING_INVALID"),
    ],
)
async def test_provider_validates_embedding_response(monkeypatch, body, code):
    provider = provider_with_response(monkeypatch, lambda _: httpx.Response(200, text=body))
    with pytest.raises(PermanentTaskError) as error:
        await provider.embed("question")
    assert error.value.code == code
