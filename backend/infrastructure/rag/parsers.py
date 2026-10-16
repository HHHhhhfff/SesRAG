import asyncio
import json
import sys
from dataclasses import asdict
from pathlib import Path
from zipfile import ZipFile

from backend.application.errors import PermanentTaskError, TransientTaskError
from backend.domain.ports import ParsedDocument, SourceLocation
from backend.infrastructure.storage.local_files import LocalFileStorage

MAX_SOURCE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_LOCATIONS = 100_000
MAX_OUTPUT_BYTES = 64 * 1024 * 1024
MAX_ZIP_ENTRIES = 100_000
PARSE_TIMEOUT = 30


def parse_document(path: Path, extension: str) -> ParsedDocument:
    """纯文本模式：行/段落位置与后续 chunk 字符切分范围分别保存。"""
    extension = extension.lower()
    if extension not in {".txt", ".md", ".docx"}:
        raise ValueError("unsupported document extension")
    try:
        if path.stat().st_size > MAX_SOURCE_BYTES:
            raise PermanentTaskError("DOCUMENT_PARSE_LIMIT", "文档解析超过资源限制")
        if extension in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8")
            parts = text.split("\n")
            kind = "line"
        else:
            from docx import Document

            # 解压前限制总展开大小，避免压缩炸弹占满内存。
            with ZipFile(path) as archive:
                if (
                    len(archive.infolist()) > MAX_ZIP_ENTRIES
                    or sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES
                ):
                    raise PermanentTaskError("DOCUMENT_PARSE_LIMIT", "文档解析超过资源限制")
            parts = [p.text for p in Document(str(path)).paragraphs]
            text = "\n".join(parts)
            kind = "paragraph"
        if text == "":
            raise PermanentTaskError("EMPTY_DOCUMENT", "文档没有可索引文本")
        if len(text) > MAX_SOURCE_BYTES or text.count("\n") >= MAX_LOCATIONS:
            raise PermanentTaskError("DOCUMENT_PARSE_LIMIT", "文档解析超过资源限制")
        offset = 0
        locations = []
        for number, part in enumerate(parts, 1):
            locations.append(SourceLocation(kind, number, offset, offset + len(part)))
            offset += len(part) + 1
        return ParsedDocument(path.name, text, tuple(locations))
    except PermanentTaskError:
        raise
    except FileNotFoundError:
        raise PermanentTaskError("DOCUMENT_SOURCE_MISSING", "文档源文件不存在") from None
    except UnicodeError:
        raise PermanentTaskError("DOCUMENT_ENCODING_INVALID", "文档必须使用 UTF-8 编码") from None
    except OSError:
        raise TransientTaskError("DOCUMENT_STORAGE_FAILED", "文档存储暂不可用") from None
    except Exception:
        raise PermanentTaskError("DOCUMENT_PARSE_FAILED", "文档解析失败") from None


class SubprocessParser:
    """解析进程在超时/取消后强制回收；不留下不可取消的后台线程。"""

    def __init__(self, root: Path):
        self.storage = LocalFileStorage(root)

    async def parse(self, storage_key: str, extension: str) -> ParsedDocument:
        try:
            path = self.storage.resolve(storage_key)
        except ValueError:
            raise PermanentTaskError("DOCUMENT_SOURCE_INVALID", "文档源标识无效") from None
        try:
            process = await asyncio.create_subprocess_exec(
                sys.executable,
                "-m",
                "backend.infrastructure.rag.parser_process",
                str(path),
                extension,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.DEVNULL,
            )
        except OSError:
            raise TransientTaskError(
                "DOCUMENT_PARSER_UNAVAILABLE", "文档解析服务暂不可用"
            ) from None
        try:
            output = bytearray()
            async with asyncio.timeout(PARSE_TIMEOUT):
                assert process.stdout is not None
                while True:
                    chunk = await process.stdout.read(65536)
                    if not chunk:
                        break
                    output.extend(chunk)
                    if len(output) > MAX_OUTPUT_BYTES:
                        raise PermanentTaskError("DOCUMENT_PARSE_LIMIT", "文档解析超过资源限制")
                await process.wait()
        except TimeoutError:
            raise PermanentTaskError("DOCUMENT_PARSE_TIMEOUT", "文档解析超时") from None
        finally:
            if process.returncode is None:
                process.kill()
                await process.wait()
        try:
            result = json.loads(bytes(output))
            if "error" in result:
                error_class = TransientTaskError if result.get("retryable") else PermanentTaskError
                raise error_class(result["error"], "文档解析失败")
            if process.returncode != 0:
                raise ValueError()
            return ParsedDocument(
                result["source_name"],
                result["text"],
                tuple(SourceLocation(**item) for item in result["locations"]),
            )
        except (ValueError, KeyError, TypeError):
            raise PermanentTaskError("DOCUMENT_PARSE_FAILED", "文档解析失败") from None


def parser_result(path: Path, extension: str) -> dict:
    return asdict(parse_document(path, extension))
